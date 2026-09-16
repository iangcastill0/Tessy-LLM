"""Shared fixtures.

Unit tests for parsing and indexing are pure-Python and run anywhere. Tests that
genuinely exercise OCR are marked `requires_tesseract` and skip cleanly when the
binary is absent, so contributors without a local build still get a useful run.
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent))  # make `synthetic` importable

from synthetic import LicenceSpec, render_licence  # noqa: E402


def _tesseract_available() -> bool:
    import os

    if os.environ.get("TESSERACT_BIN"):
        return Path(os.environ["TESSERACT_BIN"]).is_file()
    if shutil.which("tesseract") is None:
        return False
    try:
        from tessy.ocr import available_languages

        return "eng" in available_languages()
    except Exception:
        return False


def _display_available() -> bool:
    """True when a Tk window can actually be created here."""
    import os

    if not os.environ.get("DISPLAY") and sys.platform.startswith("linux"):
        return False
    try:
        import tkinter

        root = tkinter.Tk()
        root.destroy()
        return True
    except Exception:
        return False


HAVE_TESSERACT = _tesseract_available()
HAVE_DISPLAY = _display_available()

requires_tesseract = pytest.mark.skipif(
    not HAVE_TESSERACT,
    reason="tesseract with eng traineddata not available; run scripts/build_tesseract.sh",
)


def pytest_collection_modifyitems(config, items):
    """Auto-skip anything marked requires_tesseract when the binary is missing."""
    if not HAVE_TESSERACT:
        skip = pytest.mark.skip(reason="tesseract not available")
        for item in items:
            if "requires_tesseract" in item.keywords:
                item.add_marker(skip)

    if not HAVE_DISPLAY:
        skip = pytest.mark.skip(reason="no tkinter/display; run under xvfb-run")
        for item in items:
            if "requires_display" in item.keywords:
                item.add_marker(skip)


@pytest.fixture
def default_spec() -> LicenceSpec:
    return LicenceSpec()


@pytest.fixture
def texas_spec() -> LicenceSpec:
    return LicenceSpec(
        state="TEXAS",
        licence_no="T4459981",
        last_name="RIVERA",
        first_name="MARIA L",
        address_1="118 LAMAR BLVD",
        address_2="AUSTIN, TX 78701",
        dob="02/14/1985",
        expiry="02/14/2029",
        issued="02/14/2025",
        sex="F",
        eyes="HAZ",
    )


@pytest.fixture
def licence_image(tmp_path: Path, default_spec: LicenceSpec) -> Path:
    return render_licence(tmp_path / "licence.png", default_spec)


@pytest.fixture
def workbook_with_embedded_images(tmp_path: Path, default_spec, texas_spec) -> Path:
    """A workbook shaped like a real case file: text columns plus embedded scans."""
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Licences"
    ws.append(["Case Ref", "Subject Notes", "Licence Image"])

    for offset, spec in enumerate((default_spec, texas_spec)):
        row = offset + 2
        image_path = render_licence(tmp_path / f"embedded_{row}.png", spec)
        ws.cell(row=row, column=1, value=f"CASE-{2000 + row}")
        ws.cell(row=row, column=2, value=f"Interview note for row {row}")
        picture = XLImage(image_path)
        picture.width, picture.height = 300, 190
        ws.add_image(picture, f"C{row}")
        ws.row_dimensions[row].height = 150

    path = tmp_path / "case.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def workbook_with_path_column(tmp_path: Path, default_spec) -> Path:
    """A workbook that references image files by path instead of embedding them."""
    import openpyxl

    scans = tmp_path / "scans"
    image_path = render_licence(scans / "subject_a.png", default_spec)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Subjects"
    ws.append(["Case Ref", "Notes", "Scan Path"])
    ws.append(["CASE-3001", "Referred by field office", str(image_path)])
    # Second row uses a path relative to the spreadsheet.
    ws.append(["CASE-3002", "Relative path form", "scans/subject_a.png"])

    path = tmp_path / "paths.xlsx"
    wb.save(path)
    return path
