"""Tests for reading case spreadsheets. No tesseract needed."""

from __future__ import annotations

import openpyxl
import pytest

from tessy.ingest import SpreadsheetError, load_records


class TestEmbeddedImages:
    def test_reads_one_record_per_row(self, workbook_with_embedded_images, tmp_path):
        records = load_records(workbook_with_embedded_images, tmp_path / "work")
        assert [r.row for r in records] == [2, 3]

    def test_each_row_keeps_its_own_image(self, workbook_with_embedded_images, tmp_path):
        records = load_records(workbook_with_embedded_images, tmp_path / "work")
        for record in records:
            assert len(record.images) == 1
            assert record.images[0].origin == "embedded"
            assert record.images[0].path.is_file()
            assert record.images[0].path.stat().st_size > 0

    def test_images_are_anchored_to_the_right_row(self, workbook_with_embedded_images, tmp_path):
        records = load_records(workbook_with_embedded_images, tmp_path / "work")
        for record in records:
            assert record.images[0].row == record.row

    def test_text_columns_are_captured_by_header_name(
        self, workbook_with_embedded_images, tmp_path
    ):
        records = load_records(workbook_with_embedded_images, tmp_path / "work")
        assert records[0].text["Case Ref"] == "CASE-2002"
        assert "Interview note" in records[0].text["Subject Notes"]

    def test_combined_text_includes_headers(self, workbook_with_embedded_images, tmp_path):
        records = load_records(workbook_with_embedded_images, tmp_path / "work")
        assert "Case Ref: CASE-2002" in records[0].combined_text


class TestPathColumns:
    def test_absolute_and_relative_paths_both_resolve(self, workbook_with_path_column, tmp_path):
        records = load_records(workbook_with_path_column, tmp_path / "work")
        assert len(records) == 2
        for record in records:
            assert len(record.images) == 1
            assert record.images[0].origin == "path"
            assert record.images[0].path.is_file()

    def test_path_cell_is_not_also_stored_as_text(self, workbook_with_path_column, tmp_path):
        records = load_records(workbook_with_path_column, tmp_path / "work")
        assert "Scan Path" not in records[0].text
        assert records[0].text["Case Ref"] == "CASE-3001"

    def test_nonexistent_path_is_kept_as_text_not_dropped(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Case Ref", "Scan Path"])
        ws.append(["CASE-4001", "scans/missing_file.png"])
        path = tmp_path / "missing.xlsx"
        wb.save(path)

        records = load_records(path, tmp_path / "work")
        assert records[0].images == []
        # Must not be silently discarded - the operator needs to see it.
        assert records[0].text["Scan Path"] == "scans/missing_file.png"


class TestEdgeCases:
    def test_blank_rows_are_skipped(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Case Ref", "Notes"])
        ws.append(["CASE-1", "first"])
        ws.append([None, None])
        ws.append(["CASE-2", "second"])
        path = tmp_path / "gaps.xlsx"
        wb.save(path)

        assert [r.row for r in load_records(path, tmp_path / "work")] == [2, 4]

    def test_custom_header_row(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Case file export", None])
        ws.append(["Case Ref", "Notes"])
        ws.append(["CASE-1", "first"])
        path = tmp_path / "offset.xlsx"
        wb.save(path)

        records = load_records(path, tmp_path / "work", header_row=2)
        assert records[0].text["Case Ref"] == "CASE-1"

    def test_sheet_filter(self, tmp_path):
        wb = openpyxl.Workbook()
        first = wb.active
        first.title = "Wanted"
        first.append(["Case Ref"])
        first.append(["CASE-1"])
        second = wb.create_sheet("Ignored")
        second.append(["Case Ref"])
        second.append(["CASE-2"])
        path = tmp_path / "multi.xlsx"
        wb.save(path)

        records = load_records(path, tmp_path / "work", sheet="Wanted")
        assert {r.sheet for r in records} == {"Wanted"}

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(SpreadsheetError, match="No such spreadsheet"):
            load_records(tmp_path / "nope.xlsx", tmp_path / "work")

    def test_non_xlsx_raises_helpfully(self, tmp_path):
        bogus = tmp_path / "legacy.xlsx"
        bogus.write_text("this is not a zip archive")
        with pytest.raises(SpreadsheetError, match="not a readable .xlsx"):
            load_records(bogus, tmp_path / "work")

    def test_headerless_columns_fall_back_to_letters(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([None, None])
        ws.append(["value-a", "value-b"])
        path = tmp_path / "noheader.xlsx"
        wb.save(path)

        records = load_records(path, tmp_path / "work")
        assert records[0].text == {"A": "value-a", "B": "value-b"}
